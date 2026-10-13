from pathlib import Path
import json
import shutil
from threading import Event
import numpy as np
import pytest
from windstudy.pool import load
from windstudy.experiments import make_plan,Model,run_experiments,statistics
from windstudy.excel import export_report

ROOT=Path(__file__).resolve().parents[1]

def test_full_plan_coverage_and_nesting():
    pool=load(ROOT/'data/pool.sqlite')
    plan=make_plan(pool,512,100,42,Event(),lambda *x:None)
    ids={i for e in plan['plans'] for i in e['ids']}
    assert sum(plan['frequency'].values())==51200
    assert len(ids)==plan['coverage'][-1]==len(plan['cells'])
    assert all(len(set(e['ids']))==512 for e in plan['plans'])
    small=make_plan(pool,256,2,42,Event(),lambda *x:None)
    assert small['plans'][0]['ids']==plan['plans'][0]['ids'][:256]
    assert small['plans'][1]['ids']==plan['plans'][1]['ids'][:256]

def test_statistics():
    s=statistics([1,2,3]);assert s['mean']==s['median']==2
    assert s['std_sample']==1 and s['variance_sample']==1
    assert statistics([.5])['std_sample'] is None

def test_real_cache_and_excel(tmp_path):
    pool=load(ROOT/'data/pool.sqlite')
    wanted={'M2:r322:c437','M2:r306:c473','M2:r309:c384'}
    cells=[c for c in pool.cells if c[0] in wanted]
    plan={'n':2,'repetitions':2,'base_seed':42,'pool_size':len(pool.cells),'pool_checksum':pool.checksum,
          'plans':[{'experiment':1,'seed':42,'ids':[cells[0][0],cells[1][0]]},
                   {'experiment':2,'seed':43,'ids':[cells[1][0],cells[2][0]]}],
          'cells':cells,'frequency':{cells[0][0]:1,cells[1][0]:2,cells[2][0]:1},'coverage':[2,3]}
    folder=tmp_path/'cache';folder.mkdir();shutil.copy(ROOT/'cache/meteo_cache.sqlite',folder)
    def offline(*args):raise AssertionError('Unexpected network')
    from windstudy.catalog import enabled_models,load_catalog
    models=enabled_models(load_catalog(ROOT/'data/catalog.json'))
    report=run_experiments(plan,2023,50,models,folder,tmp_path/'out',Event(),lambda *x:None,offline)
    assert report['status']=='complete',report['errors']
    lookup={c['id']:c for c in report['cells']}
    expected=np.mean([lookup[c]['results'][0]['cf'] for c in plan['plans'][0]['ids']])
    assert report['means'][0][0]==pytest.approx(expected)
    assert report['statistics']['means'][0]['count']==2
    path=tmp_path/'report.xlsx';export_report(plan,report,path)
    import openpyxl
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    assert wb['Средние']['C2'].value==pytest.approx(expected)
    assert wb['Состав 1'].max_row is None or wb['Состав 1'].max_row==5
    assert len(list(wb['Состав 1'].values))==5
    assert wb['Модель power_1']['G2'].value==pytest.approx(lookup[cells[0][0]]['results'][0]['energy_kwh'])
    wb.close()

def test_cancel_keeps_no_fake_results(tmp_path):
    e=Event();e.set()
    report=run_experiments({'cells':[]},2023,50,[Model(1)],tmp_path,tmp_path/'out',e,lambda *x:None)
    assert report['status']=='cancelled' and 'means' not in report

def test_height_and_zero_policy():
    from windstudy.experiments import wind_at_height
    v,ok=wind_at_height([[2,4]],100)
    assert ok.all() and v[0]==pytest.approx(4*2**(np.log(2)/np.log(5)))
    v,ok=wind_at_height([[0,0],[np.nan,4]],50)
    assert ok.all() and list(v)==[0,4]
    _,ok=wind_at_height([[0,4]],100)
    assert not ok.all()

def test_catalog_and_safe_formulas():
    from windstudy.catalog import load_catalog,enabled_models,expression_value,CatalogModel
    data=load_catalog(ROOT/'data/catalog.json');models=enabled_models(data)
    assert len(models)==3 and len(data['curves'])==3
    for n,m in enumerate(models,1):
        assert m.power([7])[0]==pytest.approx((7**n-3**n)/(12.5**n-3**n))
    with pytest.raises(ValueError):expression_value('__import__("os").system("x")',{},np.array([5]))
    curve=data['curves'][0].copy();curve.update(nominal_kw=100,below_table='zero',above_table='hold')
    real=CatalogModel('test','Тест',100,curve);real.validate()
    assert real.power([2,3,10,30]).tolist()==pytest.approx([0,.0101,1.0106,1])
    assert real.power([3.5])[0]==pytest.approx((1.01+3.14)/2/100)
    assert data['curves'][0]['nominal_kw'] is None
